"""Inlezen en normaliseren van de maandelijkse 'VMs export' Excel-bestanden."""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

import openpyxl

# Interne kolomnaam -> regex op de (genormaliseerde) kopregel.
# De exports wisselen soms van kolomnaam (bijv. 'Total_Memory' vs 'Total Memory GB').
COLUMNS = {
    "servername": r"^server ?name$",
    "os": r"^os$",
    "klant": r"^klant$",
    "cluster": r"^cluster$",
    "cpu": r"cpu",
    "memory": r"memory",
    "tier1_single": r"tier ?1.?single",
    "tier1_dual": r"tier ?1.?dual",
    "tier3_dual": r"tier ?3.?dual",
    "mgm": r"^mgm",
    "storage": r"storage",
}
NUMERIC = ["cpu", "memory", "tier1_single", "tier1_dual", "tier3_dual", "mgm", "storage"]

# Strict OOXML (zoals de export van augustus) kan openpyxl niet lezen; we zetten
# de namespaces om naar 'transitional'.
STRICT_NS = {
    b"http://purl.oclc.org/ooxml/spreadsheetml/main": b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/docPropsVTypes": b"http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
    b"http://purl.oclc.org/ooxml/officeDocument/extendedProperties": b"http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
}
STRICT_REL = re.compile(rb"http://purl\.oclc\.org/ooxml/officeDocument/relationships/")
TRANSITIONAL_REL = b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


class ExportError(ValueError):
    pass


@dataclass
class Export:
    sheet: str
    rows: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def to_number(value) -> float:
    """'225.00', '1.234,50', 8, None -> getal. Vervangt het handmatige 'zoek .00 en vervang'."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        # Het laatste scheidingsteken is het decimaalteken.
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _strict_to_transitional(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                for old, new in STRICT_NS.items():
                    content = content.replace(old, new)
                content = STRICT_REL.sub(TRANSITIONAL_REL, content)
                content = content.replace(b' conformance="strict"', b"")
            dst.writestr(item, content)
    return out.getvalue()


def _open(data: bytes):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if wb.sheetnames:
            return wb
    except Exception:
        pass
    try:
        wb = openpyxl.load_workbook(io.BytesIO(_strict_to_transitional(data)), read_only=True, data_only=True)
    except Exception as e:
        raise ExportError(f"Kan Excel-bestand niet openen: {e}") from e
    if not wb.sheetnames:
        raise ExportError("Excel-bestand bevat geen werkbladen")
    return wb


def _norm(header) -> str:
    return re.sub(r"[\s_]+", " ", str(header or "")).strip().lower().replace("'", "")


def parse(data: bytes) -> Export:
    wb = _open(data)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        raise ExportError("Werkblad is leeg")

    index: dict[str, int] = {}
    for key, pattern in COLUMNS.items():
        for i, h in enumerate(header):
            if i not in index.values() and re.search(pattern, _norm(h)):
                index[key] = i
                break
    missing = [k for k in ("servername", "klant") if k not in index]
    if missing:
        raise ExportError(f"Kolom(men) niet gevonden: {', '.join(missing)}. Kopregel: {header}")

    export = Export(sheet=ws.title)
    for k in NUMERIC:
        if k not in index:
            export.warnings.append(f"Kolom '{k}' niet gevonden, telt als 0")

    for raw in it:
        def get(k):
            i = index.get(k)
            return raw[i] if i is not None and i < len(raw) else None

        name = get("servername")
        if name is None or str(name).strip() == "":
            continue
        row = {
            "servername": str(name).strip(),
            "os": str(get("os") or "").strip(),
            "klant": str(get("klant") or "").strip(),
            "cluster": str(get("cluster") or "").strip(),
        }
        for k in NUMERIC:
            row[k] = to_number(get(k))
        export.rows.append(row)
    wb.close()
    return export


def segments(path: str) -> list[str]:
    return [s.strip().lower() for s in path.split("/") if s.strip()]


def matches(klant: str, pattern: str) -> bool:
    """Klantfilter op mapniveau: 'Stadsbeheer' of 'Stadsbeheer/Almere' moet als hele
    map(pen) in het pad voorkomen. Zo pakt 'GGZ' niet per ongeluk ook 'GGZ GV' mee."""
    pat = segments(pattern)
    path = segments(klant)
    if not pat:
        return False
    n = len(pat)
    return any(path[i : i + n] == pat for i in range(len(path) - n + 1))
